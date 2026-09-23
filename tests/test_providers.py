"""数据层离线测试：Provider fallback / 事件去重 / 股权穿透 / CSV、XLSX 往返。

全部 mock 掉网络（requests/feedparser），不发起真实请求。
运行：python -m pytest tests/test_providers.py -q
"""
import contextlib
import csv
import gc
import os
import sys
import tempfile
import time
import unittest
from unittest import mock

import yaml

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for p in (ROOT, os.path.join(ROOT, "src")):  # 兼容 src.* 与顶层 providers.* 两种导入
    if p not in sys.path:
        sys.path.insert(0, p)

from src.core import db, penetration as pen, sync as sync_mod  # noqa: E402
from src.core import settings  # noqa: E402
from src.core.models import Enterprise, Event  # noqa: E402

# 已知问题：src/core/paths.py 的 app_root() 在开发模式下返回 src/ 而非项目根，
# settings.config() 会找不到根目录的 config.yaml。此处把真实配置预热进缓存绕开，
# 该 bug 需在 paths.py 层面另行修复（见任务报告）。
with open(os.path.join(ROOT, "config.yaml"), "r", encoding="utf-8") as _f:
    settings._config = yaml.safe_load(_f) or {}


class FakeResp:
    """模拟 requests.Response。"""

    def __init__(self, text="", status_code=200, json_data=None):
        self.text = text
        self.status_code = status_code
        self._json = json_data

    def raise_for_status(self):
        if self.status_code >= 400:
            import requests
            raise requests.HTTPError(f"HTTP {self.status_code}")

    def json(self):
        if self._json is None:
            raise ValueError("no json")
        return self._json


class BaseDBTest(unittest.TestCase):
    """把 SQLite 重定向到临时文件，避免污染项目 data 目录。"""

    def setUp(self):
        # ignore_cleanup_errors：Windows 下 WAL 句柄释放有延迟，残留临时文件交给系统回收
        self._tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        patcher = mock.patch.object(db, "db_path",
                                    return_value=os.path.join(self._tmp.name, "test.db"))
        patcher.start()
        self._db_patcher = patcher
        db.init_db()

    def tearDown(self):
        self._db_patcher.stop()
        gc.collect()  # 确认 sqlite 连接随引用计数释放，尽量让临时目录可清理
        self._tmp.cleanup()


# ---------- run_with_fallback ----------
class FallbackTest(unittest.TestCase):
    def test_switches_to_second_provider_on_exception(self):
        from src.providers import base as pbase
        ent = Enterprise(id=1, name="测试公司")
        bad = mock.Mock()
        bad.name = "bad_source"
        bad.fetch.side_effect = RuntimeError("第一个源挂了")
        good = mock.Mock()
        good.name = "good_source"
        good.fetch.return_value = [Event(category="news", title="t", source="good_source")]
        with mock.patch.object(pbase, "load_providers", return_value=[bad, good]):
            events, source = pbase.run_with_fallback("news", ent)
        self.assertEqual(source, "good_source")
        self.assertEqual(len(events), 1)
        good.fetch.assert_called_once_with(ent)

    def test_all_fail_returns_error_string(self):
        from src.providers import base as pbase
        bad = mock.Mock()
        bad.name = "bad_source"
        bad.fetch.side_effect = RuntimeError("网络错误")
        with mock.patch.object(pbase, "load_providers", return_value=[bad]):
            events, source = pbase.run_with_fallback("news", Enterprise(id=1, name="测试公司"))
        self.assertEqual(events, [])
        self.assertIn("bad_source", source)
        self.assertIn("网络错误", source)

    def test_load_providers_from_config(self):
        """config.yaml 声明的类路径可被实际加载（不发请求）。"""
        from src.providers import base as pbase
        ps = pbase.load_providers("news")
        self.assertGreaterEqual(len(ps), 2)
        self.assertIn("baidu_news", {p.name for p in ps})


# ---------- db.insert_event 去重 ----------
class DbDedupTest(BaseDBTest):
    def test_insert_event_dedup(self):
        ev = Event(enterprise_id=1, category="news", title="同一条新闻", source="baidu_news")
        self.assertTrue(db.insert_event(ev))
        self.assertFalse(db.insert_event(ev))  # 唯一键命中，去重
        ev2 = Event(enterprise_id=1, category="news", title="同一条新闻", source="sogou_news")
        self.assertTrue(db.insert_event(ev2))  # 不同 source 不算重复
        self.assertEqual(len(db.list_events(enterprise_id=1)), 2)


# ---------- run_sync ----------
class RunSyncTest(BaseDBTest):
    def test_run_sync_insert_and_dedup(self):
        def fake_rwf(cat, ent):
            if cat == "news":
                return [Event(category="news", title="测试新闻", source="baidu_news")], "baidu_news"
            return [], f"{cat}: empty"

        ent = Enterprise(id=1, name="测试公司")
        with mock.patch.object(sync_mod, "run_with_fallback", side_effect=fake_rwf):
            res = sync_mod.run_sync(ent)
        self.assertEqual(res["news"]["inserted"], 1)
        self.assertEqual(res["news"]["source"], "baidu_news")
        self.assertEqual(res["biz_change"]["inserted"], 0)
        with mock.patch.object(sync_mod, "run_with_fallback", side_effect=fake_rwf):
            res2 = sync_mod.run_sync(ent)  # 第二次同步全部去重
        self.assertEqual(res2["news"]["inserted"], 0)


# ---------- 股权穿透 ----------
TREE = {
    "甲公司": {"up": [{"name": "乙控股", "ratio": 60}],
              "down": [{"name": "丙子公司", "ratio": 100}]},
    "乙控股": {"up": [{"name": "张三", "ratio": 100}], "down": []},
    "丙子公司": {"up": [{"name": "甲公司", "ratio": 100}],
               "down": [{"name": "丁孙公司", "ratio": 51},
                        {"name": "甲公司", "ratio": 100}]},
    "张三": {"up": [], "down": []},
    "丁孙公司": {"up": [], "down": []},
}


def _shareholders(name):
    return list(TREE.get(name, {}).get("up", []))


def _investments(name):
    return list(TREE.get(name, {}).get("down", []))


class PenetrationTest(BaseDBTest):
    def _patch(self):
        stack = contextlib.ExitStack()
        stack.enter_context(mock.patch.object(pen, "fetch_shareholders", _shareholders))
        stack.enter_context(mock.patch.object(pen, "fetch_investments", _investments))
        return stack

    def test_levels_relations_and_dedup(self):
        with self._patch():
            out = pen.penetrate(Enterprise(id=1, name="甲公司"), up_levels=2, down_levels=2)
        by_name = {e.name: e for e in out}
        self.assertEqual({"乙控股", "张三", "丙子公司", "丁孙公司"}, set(by_name))
        # 循环股权（丙 -> 甲）被去重，甲公司不重复出现
        self.assertEqual([e.name for e in out].count("甲公司"), 0)
        # 层级 / 方向 / 关系
        self.assertEqual(by_name["乙控股"].pen_level, 1)
        self.assertEqual(by_name["乙控股"].direction, "up")
        self.assertEqual(by_name["乙控股"].relation, "股东(60%)")
        self.assertEqual(by_name["张三"].pen_level, 2)
        self.assertEqual(by_name["丙子公司"].relation, "对外投资(100%)")
        self.assertEqual(by_name["丁孙公司"].relation, "对外投资(51%)")
        self.assertEqual(by_name["丁孙公司"].direction, "down")

    def test_level_limit(self):
        with self._patch():
            out = pen.penetrate(Enterprise(id=1, name="甲公司"), up_levels=1, down_levels=1)
        names = {e.name for e in out}
        self.assertEqual(names, {"乙控股", "丙子公司"})  # 第二层被层数限制截断

    def test_fetch_exception_tolerated(self):
        def boom(name):
            raise RuntimeError("爱企查被反爬")
        with mock.patch.object(pen, "fetch_shareholders", boom), \
             mock.patch.object(pen, "fetch_investments", _investments):
            out = pen.penetrate(Enterprise(id=1, name="甲公司"), up_levels=2, down_levels=2)
        self.assertEqual({e.name for e in out}, {"丙子公司", "丁孙公司"})

    def test_existing_enterprise_skipped(self):
        db.upsert_enterprise("乙控股")
        with self._patch():
            out = pen.penetrate(Enterprise(id=1, name="甲公司"), up_levels=2, down_levels=2)
        self.assertNotIn("乙控股", {e.name for e in out})

    def test_save_related(self):
        with self._patch():
            out = pen.penetrate(Enterprise(id=1, name="甲公司"), up_levels=2, down_levels=2)
        saved = pen.save_related(1, out, ["乙控股", "未勾选的"])
        self.assertEqual(len(saved), 1)
        rec = db.get_enterprise(saved[0])
        self.assertEqual(rec.name, "乙控股")
        self.assertEqual(rec.parent_id, 1)
        self.assertEqual(rec.is_related, 1)
        self.assertEqual(rec.relation, "股东(60%)")
        self.assertEqual(rec.pen_level, 1)
        self.assertEqual(rec.direction, "up")


# ---------- CSV / XLSX 导入导出 ----------
class ImportExportTest(BaseDBTest):
    CSV_CONTENT = ("企业名称,统一社会信用代码,省份,城市,授信备注\n"
                   "测试一号公司,91330100MA27AAAA1A,浙江省,杭州市,授信500万\n"
                   "测试二号公司,,,宁波市,\n")

    def test_csv_roundtrip(self):
        src = os.path.join(self._tmp.name, "in.csv")
        with open(src, "w", encoding="utf-8-sig", newline="") as f:
            f.write(self.CSV_CONTENT)
        self.assertEqual(sync_mod.import_enterprises_from_csv(src), 2)
        dst = os.path.join(self._tmp.name, "out.csv")
        self.assertEqual(sync_mod.export_enterprises_to_csv(dst), 2)
        with open(dst, "r", encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        self.assertEqual(rows[0], ["企业名称", "统一社会信用代码", "省份", "城市", "授信备注"])
        self.assertEqual(rows[1], ["测试一号公司", "91330100MA27AAAA1A", "浙江省", "杭州市", "授信500万"])
        self.assertEqual(rows[2][:2], ["测试二号公司", ""])

    def test_csv_gbk_import(self):
        src = os.path.join(self._tmp.name, "gbk.csv")
        with open(src, "w", encoding="gbk", newline="") as f:
            f.write(self.CSV_CONTENT)
        self.assertEqual(sync_mod.import_enterprises_from_csv(src), 2)

    def test_xlsx_roundtrip(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl 未安装")
        src = os.path.join(self._tmp.name, "in.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["企业名称", "统一社会信用代码", "省份", "城市", "授信备注"])
        ws.append(["测试三号公司", "91330100MA27BBBB2B", "浙江省", "温州市", ""])
        wb.save(src)
        self.assertEqual(sync_mod.import_enterprises_from_xlsx(src), 1)
        dst = os.path.join(self._tmp.name, "out.xlsx")
        self.assertEqual(sync_mod.export_enterprises_to_xlsx(dst), 1)
        wb2 = openpyxl.load_workbook(dst, read_only=True)
        try:
            rows = list(wb2.active.iter_rows(values_only=True))
        finally:
            wb2.close()
        self.assertEqual(rows[0], ("企业名称", "统一社会信用代码", "省份", "城市", "授信备注"))
        # openpyxl 读回时空字符串单元格为 None，统一归一化后比较
        self.assertEqual(tuple("" if v is None else v for v in rows[1]),
                         ("测试三号公司", "91330100MA27BBBB2B", "浙江省", "温州市", ""))


# ---------- Provider 离线解析 ----------
BAIDU_HTML = """<html><body>
<div class="result-op"><h3><a href="http://example.com/1">测试公司 中标大单</a></h3>
<div class="c-abstract">测试公司中标某重点项目</div>
<span class="c-color-gray">新华网 3小时前</span></div>
<div class="result"><h3><a href="http://example.com/2">测试公司被处罚</a></h3>
<div class="c-abstract">涉嫌违规被罚款</div>
<span class="c-color-gray">央视网 2026-09-20</span></div>
</body></html>"""

SOGOU_HTML = """<html><body>
<div class="vrwrap"><h3><a href="/link?url=abc">测试公司 签署战略合作</a></h3>
<div class="news-txt"><p>测试公司与某银行签约</p></div>
<div class="news-from">中国新闻网 2026-09-19 10:00</div></div>
</body></html>"""


class BaiduNewsTest(unittest.TestCase):
    def _provider(self):
        from src.providers import news_baidu as m
        return m, m.BaiduNewsProvider({})

    def test_parse(self):
        m, p = self._provider()
        with mock.patch.object(p, "http_get", return_value=FakeResp(BAIDU_HTML)):
            events = p.fetch(Enterprise(id=1, name="测试公司"))
        self.assertEqual(len(events), 2)
        self.assertEqual(events[0].source, "baidu_news")
        self.assertEqual(events[0].category, "news")
        self.assertEqual(events[0].sentiment, "positive")   # 中标 -> 正面
        self.assertEqual(events[0].occurred_at, time.strftime("%Y-%m-%d"))
        self.assertEqual(events[1].sentiment, "negative")   # 处罚 -> 负面
        self.assertEqual(events[1].severity, "medium")
        self.assertEqual(events[1].occurred_at, "2026-09-20")
        self.assertTrue(all(e.relevance == "high" for e in events))

    def test_anti_bot_raises(self):
        _, p = self._provider()
        with mock.patch.object(p, "http_get",
                               return_value=FakeResp("<html>百度安全验证</html>", status_code=200)):
            with self.assertRaises(RuntimeError):
                p.fetch(Enterprise(id=1, name="测试公司"))


class SogouNewsTest(unittest.TestCase):
    def test_parse(self):
        from src.providers import news_sogou as m
        p = m.SogouNewsProvider({})
        with mock.patch.object(p, "http_get", return_value=FakeResp(SOGOU_HTML)):
            events = p.fetch(Enterprise(id=1, name="测试公司"))
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0].source, "sogou_news")
        self.assertEqual(events[0].url, "https://news.sogou.com/link?url=abc")
        self.assertEqual(events[0].occurred_at, "2026-09-19")


class RssTest(unittest.TestCase):
    def test_keyword_filter(self):
        from src.providers import news_rss as m
        fake = mock.Mock(entries=[
            {"title": "测试公司 发布新产品", "summary": "详见全文", "link": "http://x/1",
             "published_parsed": time.struct_time((2026, 9, 20, 0, 0, 0, 0, 0, 0))},
            {"title": "完全无关的国际新闻", "summary": "other", "link": "http://x/2"},
        ])
        p = m.RssProvider({"feeds": ["http://fake/rss"]})
        with mock.patch.object(m.feedparser, "parse", return_value=fake) as fp:
            events = p.fetch(Enterprise(id=1, name="测试公司"))
        fp.assert_called_once()
        self.assertEqual(len(events), 1)  # 关键词过滤只保留命中条目
        self.assertEqual(events[0].title, "测试公司 发布新产品")
        self.assertEqual(events[0].source, "rss")
        self.assertEqual(events[0].occurred_at, "2026-09-20")


class IsteroTest(unittest.TestCase):
    def test_token_empty_raises(self):
        from src.providers import news_istero as m
        p = m.IsteroNewsProvider({"base_url": "http://fake/api", "token": ""})
        with self.assertRaises(RuntimeError):
            p.fetch(Enterprise(id=1, name="测试公司"))

    def test_fetch_with_token(self):
        from src.providers import news_istero as m
        data = {"data": {"list": [{"title": "测试公司 登上热榜"},
                                  {"title": "无关热搜条目"},
                                  {"title": "测试公司控股重组"}]}}
        p = m.IsteroNewsProvider({"base_url": "http://fake/api", "token": "abc"})
        with mock.patch.object(p, "http_get", return_value=FakeResp(json_data=data)):
            events = p.fetch(Enterprise(id=1, name="测试公司"))
        self.assertEqual(len(events), 3)
        hit = [e for e in events if e.relevance == "high"]
        self.assertEqual(len(hit), 2)  # 标题含企业名 relevance 提高
        self.assertTrue(all(e.source == "istero_pengpai" for e in events))
        miss = [e for e in events if e.relevance == "low"]
        self.assertTrue(all(e.severity == "low" and e.sentiment == "neutral" for e in miss))


class NewsApiTest(unittest.TestCase):
    def test_api_key_empty_raises(self):
        from src.providers import newsapi_org as m
        p = m.NewsApiProvider({"api_key": ""})
        with self.assertRaises(RuntimeError):
            p.fetch(Enterprise(id=1, name="测试公司"))


class AiqichaTest(unittest.TestCase):
    # 样本页需超过 500 字符以通过"反爬页过短"启发式检查
    SEARCH_HTML = ('<html><body>' + '<!-- ' + 'pad ' * 100 + ' -->'
                   '<a href="https://aiqicha.baidu.com/company_detail_12345">测试公司</a>'
                   '</body></html>')

    def test_change_events(self):
        from src.providers import biz_aiqicha as m
        p = m.AiqichaProvider({})
        changes = {"data": {"list": [
            {"changeItem": "股东变更", "beforeContent": "张三", "afterContent": "李四",
             "changeDate": "2026-08-01"},
            {"changeItem": "法定代表人变更", "beforeContent": "王五", "afterContent": "赵六",
             "changeDate": "2026-08-02"},
            {"changeItem": "其他备案变更", "beforeContent": "a", "afterContent": "b"},
        ]}}
        responses = [FakeResp(self.SEARCH_HTML), FakeResp(json_data=changes)]
        with mock.patch.object(p, "http_get", side_effect=responses):
            events = p.fetch(Enterprise(id=1, name="测试公司"))
        self.assertEqual(len(events), 2)
        self.assertEqual(events[0].title, "测试公司：股东由张三变更为李四")
        self.assertEqual(events[0].category, "biz_change")
        self.assertEqual(events[0].severity, "medium")
        self.assertEqual(events[1].title, "测试公司：法定代表人由王五变更为赵六")

    def test_fetch_shareholders_mocked(self):
        from src.providers import biz_aiqicha as m
        holders = {"data": {"list": [{"subName": "李四", "ratio": "60%"},
                                     {"subName": "张三", "ratio": 0.4}]}}
        with mock.patch.object(m.AiqichaProvider, "http_get",
                               side_effect=[FakeResp(self.SEARCH_HTML),
                                            FakeResp(json_data=holders)]):
            out = m.fetch_shareholders("测试公司")
        self.assertEqual(out, [{"name": "李四", "ratio": 60.0}, {"name": "张三", "ratio": 40.0}])


class CommercialApiTest(unittest.TestCase):
    def test_params_empty_raises(self):
        from src.providers import biz_commercial as m
        p = m.CommercialApiProvider({"base_url": "", "api_key": "", "vendor": "tianyancha"})
        with self.assertRaises(RuntimeError):
            p.fetch(Enterprise(id=1, name="测试公司"))

    def test_tianyancha_mapping(self):
        from src.providers import biz_commercial as m
        p = m.CommercialApiProvider({"base_url": "http://fake/open", "api_key": "k",
                                     "vendor": "tianyancha"})
        baseinfo = {"result": {"regStatus": "存续", "legalPersonName": "王五",
                               "regCapital": "1000万人民币", "id": 123}}
        holders = {"result": {"result": [{"name": "李四", "ratio": "60%"}]}}
        with mock.patch.object(p, "http_get",
                               side_effect=[FakeResp(json_data=baseinfo), FakeResp(json_data=holders)]):
            events = p.fetch(Enterprise(id=1, name="测试公司"))
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0].title, "测试公司：登记状态存续")
        self.assertEqual(events[0].severity, "low")
        self.assertIn("王五", events[0].summary)
        self.assertIn("李四", events[0].summary)

    def test_other_vendor_not_implemented(self):
        from src.providers import biz_commercial as m
        p = m.CommercialApiProvider({"base_url": "http://fake", "api_key": "k", "vendor": "qichacha"})
        with self.assertRaises(NotImplementedError):
            p.fetch(Enterprise(id=1, name="测试公司"))


class ZxgkTest(unittest.TestCase):
    def test_hit(self):
        from src.providers import court_zxgk as m
        p = m.ZxgkProvider({})
        data = {"status": 1, "data": [{"caseCode": "(2026)浙01执123号",
                                       "execCourtName": "杭州市中级人民法院",
                                       "caseCreateTime": "2026-08-01 00:00:00"}]}
        with mock.patch.object(p, "http_post", return_value=FakeResp(json_data=data)):
            events = p.fetch(Enterprise(id=1, name="测试公司"))
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0].category, "enforcement")
        self.assertEqual(events[0].severity, "high")
        self.assertIn("被执行人", events[0].title)
        self.assertIn("(2026)浙01执123号", events[0].title)

    def test_captcha_raises(self):
        from src.providers import court_zxgk as m
        p = m.ZxgkProvider({})
        with mock.patch.object(p, "http_post",
                               return_value=FakeResp(json_data={"status": 0, "message": "验证码错误"})):
            with self.assertRaises(RuntimeError):
                p.fetch(Enterprise(id=1, name="测试公司"))


class WenshuTest(unittest.TestCase):
    def test_blocked_raises(self):
        from src.providers import court_wenshu as m
        p = m.WenshuProvider({})
        with mock.patch.object(p, "http_get", return_value=FakeResp("<html>ok</html>")), \
             mock.patch.object(p, "http_post",
                               return_value=FakeResp("<!DOCTYPE html><html>请开启JavaScript</html>")):
            with self.assertRaises(RuntimeError):
                p.fetch(Enterprise(id=1, name="测试公司"))


if __name__ == "__main__":
    unittest.main()
