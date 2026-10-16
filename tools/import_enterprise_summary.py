"""从"企业信息汇总.xlsx"主表导入监控企业（真实名单，替代演示数据）。

- 读《企业信息汇总（主表）》sheet：企业名称/统一社会信用代码/注册地址/
  股东/是否涉诉/是否被执行/经营异常/成立日期/注册资本/经营状态
- 省市解析：地址匹配省名；"沈抚示范区"归入抚顺市（用于地图定位）
- 授信备注：汇编股东/涉诉/被执行/经营异常等风险要点
- 同时软删除 seed_demo.py 生成的 12 家演示企业（历史事件保留）

运行：D:/hermes/hermes-agent/venv/Scripts/python.exe tools/import_enterprise_summary.py [xlsx路径]
"""
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from src.core import db  # noqa: E402

DEFAULT_XLSX = r"C:\Users\Administrator\Desktop\新建文件夹 (2)\企业信息汇总.xlsx"
SHEET = "企业信息汇总（主表）"

# 演示数据企业名（tools/seed_demo.py），导入真实名单前软删除
DEMO_NAMES = [
    "恒信实业集团有限公司", "中科恒泰新能源科技有限公司", "华辰置业发展有限公司",
    "金桥港口物流股份有限公司", "天晟精密制造有限公司", "蜀都食品集团有限公司",
    "蓝湾海洋生物科技有限公司", "北疆重工装备股份有限公司", "楚天医药化工有限公司",
    "闽东南电子信息技术有限公司", "中原农业发展集团有限公司", "云上大数据服务股份有限公司",
]


def parse_region(address: str) -> tuple:
    """从注册地址解析 (省份, 城市)。沈抚示范区归入抚顺市便于地图定位。"""
    addr = str(address or "")
    m = re.search(r"(.+?(?:省|自治区|北京市|天津市|上海市|重庆市))", addr)
    prov = m.group(1) if m else ""
    if "沈抚" in addr:
        return (prov or "辽宁省"), "抚顺市"
    m2 = re.search(r"(?:省|自治区)?([\u4e00-\u9fa5]{2,7}市)", addr)
    city = m2.group(1) if m2 else ""
    return prov, city


def main() -> int:
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    from openpyxl import load_workbook

    # 1) 软删除演示企业（历史事件保留可查）
    removed = 0
    for row in db.list_enterprises(include_stopped=True):
        if row.name in DEMO_NAMES:
            db.delete_enterprise(row.id)
            removed += 1
    print(f"已停止监控演示企业 {removed} 家（历史事件保留）")

    # 2) 导入真实名单
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {str(h).strip().replace("\n", ""): i for i, h in enumerate(header) if h}
    imported = 0
    for r in rows[1:]:
        if not r or not r[idx.get("企业名称", 1)]:
            continue
        name = str(r[idx["企业名称"]]).strip()
        i_code = idx.get("统一社会信用代码")
        code = str(r[i_code] or "").strip() if i_code is not None and i_code < len(r) else ""
        addr = r[idx.get("注册地址", 11)] or ""
        prov, city = parse_region(addr)
        note_parts = []
        for label, key in (("成立", "成立日期"), ("注册资本", "注册资本(万元)"),
                           ("经营状态", "经营状态"), ("股东", "股东"),
                           ("涉诉", "是否涉诉"), ("被执行", "是否被执行"),
                           ("经营异常", "经营异常")):
            i = idx.get(key)
            v = r[i] if i is not None and i < len(r) else None
            if v is not None and str(v).strip() and str(v).strip() != "—":
                note_parts.append(f"{label}：{str(v).strip()}")
        i_loan = idx.get("放款日期")
        loan_date = (str(r[i_loan]).strip()
                     if i_loan is not None and i_loan < len(r) and r[i_loan] is not None
                     else "")
        eid = db.upsert_enterprise(
            name=name, credit_code=code, province=prov, city=city,
            credit_note="；".join(note_parts), loan_date=loan_date)
        print(f"[ok] {name}  {prov}/{city}  (id={eid})")
        imported += 1
    wb.close()
    print(f"\n共导入 {imported} 家企业。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
